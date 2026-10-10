"""Standalone document extraction and OCR owned entirely by PAi Legal.

The Store package bundles Tesseract and its English language data under
resources/tesseract.  Development builds may also use an installed Tesseract.
No other PAi product or PSi Legal component is required to ingest a file.
"""
from __future__ import annotations

import csv
import email
import html
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from dataclasses import dataclass
from email import policy
from pathlib import Path
from typing import Iterable, List, Optional
from xml.etree import ElementTree

from .procutil import hidden


class ExtractionError(RuntimeError):
    pass


@dataclass(frozen=True)
class ExtractionResult:
    text: str
    method: str
    pages: int = 0
    ocr_pages: int = 0


def application_root() -> Path:
    """Directory containing packaged runtime resources."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent.parent


def find_tesseract() -> Optional[Path]:
    override = os.environ.get("PAI_LEGAL_TESSERACT")
    candidates: List[Path] = []
    if override:
        candidates.append(Path(override))
    root = application_root()
    candidates.extend([
        root / "resources" / "tesseract" / "tesseract.exe",
        root / "resources" / "tesseract" / "tesseract",
        Path(os.environ.get("ProgramFiles", "")) / "Tesseract-OCR" / "tesseract.exe",
        Path(os.environ.get("LOCALAPPDATA", "")) / "Programs" / "Tesseract-OCR" / "tesseract.exe",
    ])
    executable = shutil.which("tesseract")
    if executable:
        candidates.append(Path(executable))
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    return None


class TextExtractor:
    TEXT_TYPES = {".txt", ".md", ".log", ".json", ".xml", ".yaml", ".yml"}
    IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".gif"}

    def __init__(self, language: str = "eng", tesseract: Optional[Path] = None):
        self.language = language
        self.tesseract = Path(tesseract) if tesseract else find_tesseract()

    @property
    def ocr_available(self) -> bool:
        return bool(self.tesseract and self.tesseract.is_file())

    def extract(self, source: Path) -> ExtractionResult:
        source = Path(source)
        suffix = source.suffix.lower()
        if suffix in self.TEXT_TYPES:
            return ExtractionResult(source.read_text(encoding="utf-8", errors="replace"), "text")
        if suffix in {".html", ".htm"}:
            raw = source.read_text(encoding="utf-8", errors="replace")
            text = html.unescape(re.sub(r"<[^>]+>", " ", raw))
            return ExtractionResult(self._collapse(text), "html")
        if suffix in {".csv", ".tsv"}:
            delimiter = "\t" if suffix == ".tsv" else ","
            with source.open("r", encoding="utf-8", errors="replace", newline="") as handle:
                text = "\n".join(" | ".join(row) for row in csv.reader(handle, delimiter=delimiter))
            return ExtractionResult(text, "delimited-text")
        if suffix == ".pdf":
            return self._pdf(source)
        if suffix in self.IMAGE_TYPES:
            return ExtractionResult(self._ocr(source), "tesseract-ocr", pages=1, ocr_pages=1)
        if suffix == ".docx":
            return ExtractionResult(self._docx(source), "docx")
        if suffix == ".odt":
            return ExtractionResult(self._odt(source), "odt")
        if suffix == ".xlsx":
            return ExtractionResult(self._xlsx(source), "xlsx")
        if suffix == ".pptx":
            return ExtractionResult(self._pptx(source), "pptx")
        if suffix == ".rtf":
            return ExtractionResult(self._rtf(source), "rtf")
        if suffix == ".eml":
            return ExtractionResult(self._eml(source), "eml")
        if suffix == ".msg":
            return ExtractionResult(self._msg(source), "msg")
        raise ExtractionError(f"Unsupported file type: {suffix or 'no extension'}")

    def _ocr(self, image: Path) -> str:
        if not self.ocr_available:
            raise ExtractionError("The bundled Tesseract OCR runtime was not found")
        environment = os.environ.copy()
        tessdata = self.tesseract.parent / "tessdata"
        if tessdata.is_dir():
            environment["TESSDATA_PREFIX"] = str(tessdata)
        command = [
            str(self.tesseract), str(image), "stdout", "-l", self.language,
            "--oem", "1", "--psm", "3",
        ]
        try:
            result = subprocess.run(
                command, capture_output=True, text=True, encoding="utf-8",
                errors="replace", timeout=600, env=environment, **hidden(),
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            raise ExtractionError(f"Tesseract failed to run: {exc}") from exc
        if result.returncode:
            raise ExtractionError(result.stderr.strip() or f"Tesseract exited with code {result.returncode}")
        return result.stdout.strip()

    def _pdf(self, path: Path) -> ExtractionResult:
        try:
            import pymupdf
        except ImportError as exc:
            raise ExtractionError("The bundled PDF renderer is not available") from exc
        pieces: List[str] = []
        ocr_pages = 0
        try:
            document = pymupdf.open(str(path))
        except Exception as exc:
            raise ExtractionError(f"Cannot open PDF: {exc}") from exc
        try:
            for number, page in enumerate(document, 1):
                page_text = (page.get_text("text") or "").strip()
                # A page with only a stamp/page number is still effectively scanned.
                if len(re.sub(r"\s+", "", page_text)) < 40:
                    if not self.ocr_available:
                        raise ExtractionError(
                            f"PDF page {number} needs OCR, but bundled Tesseract was not found"
                        )
                    with tempfile.TemporaryDirectory(prefix="pai-legal-ocr-") as temp:
                        image = Path(temp) / f"page-{number:05d}.png"
                        pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2.5, 2.5), alpha=False)
                        pixmap.save(str(image))
                        page_text = self._ocr(image)
                        ocr_pages += 1
                pieces.append(f"<<<PAGE {number}>>>\n{page_text}")
            return ExtractionResult(
                "\n\n".join(pieces),
                "pdf-text+tesseract" if ocr_pages else "pdf-text",
                pages=len(document),
                ocr_pages=ocr_pages,
            )
        finally:
            document.close()

    @staticmethod
    def _docx(path: Path) -> str:
        with zipfile.ZipFile(path) as archive:
            raw = archive.read("word/document.xml")
        root = ElementTree.fromstring(raw)
        namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        paragraphs: List[str] = []
        for paragraph in root.iter(namespace + "p"):
            text = "".join(node.text or "" for node in paragraph.iter(namespace + "t"))
            if text:
                paragraphs.append(text)
        return "\n".join(paragraphs)

    @staticmethod
    def _odt(path: Path) -> str:
        with zipfile.ZipFile(path) as archive:
            raw = archive.read("content.xml")
        root = ElementTree.fromstring(raw)
        paragraphs: List[str] = []
        for element in root.iter():
            if element.tag.endswith("}p") or element.tag.endswith("}h"):
                text = "".join(element.itertext()).strip()
                if text:
                    paragraphs.append(text)
        return "\n".join(paragraphs)

    @staticmethod
    def _xlsx(path: Path) -> str:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ExtractionError("The bundled Excel extractor is not available") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        sections: List[str] = []
        try:
            for sheet in workbook.worksheets:
                sections.append(f"<<<SHEET {sheet.title}>>>")
                for row in sheet.iter_rows(values_only=True):
                    if any(value is not None for value in row):
                        sections.append(" | ".join("" if value is None else str(value) for value in row))
        finally:
            workbook.close()
        return "\n".join(sections)

    @staticmethod
    def _pptx(path: Path) -> str:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise ExtractionError("The bundled PowerPoint extractor is not available") from exc
        presentation = Presentation(str(path))
        sections: List[str] = []
        for number, slide in enumerate(presentation.slides, 1):
            sections.append(f"<<<SLIDE {number}>>>")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    sections.append(shape.text.strip())
        return "\n".join(sections)

    @staticmethod
    def _rtf(path: Path) -> str:
        try:
            from striprtf.striprtf import rtf_to_text
        except ImportError as exc:
            raise ExtractionError("The bundled RTF extractor is not available") from exc
        return rtf_to_text(path.read_text(encoding="utf-8", errors="replace"))

    @staticmethod
    def _eml(path: Path) -> str:
        message = email.message_from_bytes(path.read_bytes(), policy=policy.default)
        sections = [
            f"From: {message.get('From', '')}", f"To: {message.get('To', '')}",
            f"Cc: {message.get('Cc', '')}", f"Date: {message.get('Date', '')}",
            f"Subject: {message.get('Subject', '')}", "",
        ]
        bodies: List[str] = []
        if message.is_multipart():
            for part in message.walk():
                if part.get_content_disposition() == "attachment":
                    sections.append(f"Attachment: {part.get_filename() or 'unnamed'}")
                elif part.get_content_type() == "text/plain":
                    try:
                        bodies.append(part.get_content())
                    except Exception:
                        pass
        else:
            try:
                bodies.append(message.get_content())
            except Exception:
                bodies.append(message.get_payload(decode=True).decode(errors="replace"))
        return "\n".join(sections + bodies)

    @staticmethod
    def _msg(path: Path) -> str:
        try:
            import extract_msg
        except ImportError as exc:
            raise ExtractionError("The bundled Outlook MSG extractor is not available") from exc
        message = extract_msg.Message(str(path))
        try:
            attachments = [getattr(item, "longFilename", None) or getattr(item, "shortFilename", None) or "unnamed" for item in message.attachments]
            return "\n".join([
                f"From: {message.sender or ''}", f"To: {message.to or ''}",
                f"Cc: {message.cc or ''}", f"Date: {message.date or ''}",
                f"Subject: {message.subject or ''}",
                *(f"Attachment: {name}" for name in attachments), "", message.body or "",
            ])
        finally:
            message.close()

    @staticmethod
    def _collapse(text: str) -> str:
        return re.sub(r"[ \t]+", " ", re.sub(r"\n\s*\n\s*\n+", "\n\n", text)).strip()
